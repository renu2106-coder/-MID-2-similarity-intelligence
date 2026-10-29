import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient
from main import app
import pymupdf as fitz
import docx

def test_multiformat():
    test_pdf_path = os.path.join(os.path.dirname(__file__), "test_sample.pdf")
    test_docx_path = os.path.join(os.path.dirname(__file__), "test_sample.docx")

    # 1. Create a real sample PDF using PyMuPDF
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Cardiovascular Risk Biomarkers in Preventive Cardiology", fontsize=14)
    page1.insert_text((50, 120), "Elevated levels of high-sensitivity C-reactive protein (hs-CRP) indicate systemic vascular inflammation.", fontsize=11)
    page1.insert_text((50, 160), "Coronary artery calcium (CAC) scoring via non-contrast cardiac CT stratifies atherosclerotic plaque burden.", fontsize=11)
    
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Lipoprotein(a) and Genetic Predisposition to Myocardial Infarction", fontsize=14)
    page2.insert_text((50, 120), "Lipoprotein(a) is an independent genetically determined causal risk factor for premature atherosclerotic disease.", fontsize=11)
    doc.save(test_pdf_path)
    doc.close()
    print("Created test PDF:", test_pdf_path)

    # 2. Create a real sample DOCX using python-docx
    docx_doc = docx.Document()
    docx_doc.add_heading("Quantum Computing Algorithms in Cryptography", level=1)
    docx_doc.add_paragraph("Shor's algorithm demonstrates polynomial time factorization of large composite integers on fault-tolerant quantum hardware.")
    docx_doc.add_paragraph("Grover's algorithm provides a quadratic speedup for unstructured database search and symmetric key cryptanalysis.")
    docx_doc.save(test_docx_path)
    print("Created test DOCX:", test_docx_path)

    # 3. Test uploading both via TestClient
    with TestClient(app) as client:
        with open(test_pdf_path, "rb") as f_pdf, open(test_docx_path, "rb") as f_docx:
            response = client.post(
                "/api/documents/upload",
                files=[
                    ("files", ("test_sample.pdf", f_pdf, "application/pdf")),
                    ("files", ("test_sample.docx", f_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
                ]
            )

        assert response.status_code == 200, f"Upload failed: {response.text}"
        data = response.json()
        print("\nUpload response:", data)
        assert len(data["uploaded"]) == 2
        for item in data["uploaded"]:
            assert item["status"] == "indexed"
            print(f"Uploaded and indexed {item['filename']} with {item['chunks']} chunks.")

        # 4. Search query matching the PDF
        print("\n--- Searching for PDF concept ---")
        pdf_res = client.post("/api/search", json={
            "query": "What does coronary calcium scoring evaluate in vascular inflammation?",
            "threshold": 0.35,
            "top_k": 3
        })
        assert pdf_res.status_code == 200
        pdf_data = pdf_res.json()
        print(f"Matches for cardiac CT query: {pdf_data['total_matches']}")
        assert pdf_data["total_matches"] > 0
        top_pdf = pdf_data["results"][0]
        print(f"Top match: {top_pdf['filename']}, Page: {top_pdf['page_number']}, Coordinates: {top_pdf['bbox']}")
        assert "test_sample.pdf" in top_pdf["filename"]
        assert top_pdf["page_number"] == 1
        assert top_pdf["bbox"] is not None

        # 5. Search query matching the DOCX
        print("\n--- Searching for DOCX concept ---")
        docx_res = client.post("/api/search", json={
            "query": "How does Shor's algorithm impact integer factorization and encryption?",
            "threshold": 0.35,
            "top_k": 3
        })
        assert docx_res.status_code == 200
        docx_data = docx_res.json()
        print(f"Matches for Shor's algorithm query: {docx_data['total_matches']}")
        assert docx_data["total_matches"] > 0
        top_docx = docx_data["results"][0]
        print(f"Top match: {top_docx['filename']}, Paragraph: {top_docx['paragraph_number']}, Lines: {top_docx['line_start']}-{top_docx['line_end']}")
        assert "test_sample.docx" in top_docx["filename"]

    # Clean up test files
    try:
        os.remove(test_pdf_path)
        os.remove(test_docx_path)
    except Exception:
        pass

    print("\nMULTIFORMAT (PDF & DOCX) TESTS COMPLETED SUCCESSFULLY!")

if __name__ == "__main__":
    test_multiformat()

